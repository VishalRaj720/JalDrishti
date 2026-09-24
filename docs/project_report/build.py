"""Build the TEXMiN UDAAN fellowship project report (.docx) from the text sources.

    python docs/project_report/build.py            # build once
    python docs/project_report/build.py --pages    # build, render with Word, fill the
                                                   # contents page numbers, build again

Source files: front.json (cover, certificate, declaration, acknowledgement, abstract,
contents), ch*.txt (chapters, references, annexures), back.json (remarks, approval,
checklist) and references.json (reference library, keyed).

Markup, one block per blank-line-separated paragraph:
    # 4. METHODOLOGY / APPROACH   chapter heading, starts a new page
    ## 4.1 Requirement Analysis  section heading
    ### Activities Performed     small bold heading
    - item / -- sub-item         bullet list
    1. item                      numbered item (number kept as written)
    TABLE: caption               a table follows: |a|b| rows, first row is the header
    ![caption](path){w=15}       figure, width in cm, auto-numbered
    FLOW: A >> B >> C            centred vertical workflow with arrows
    EQ: formula                  centred formula line
    [@key]                       citation, numbered by first use
    ```  ...  ```                 code block (monospace, shaded)
    **bold**  *italic*           inline
    PAGEBREAK
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
OUT = HERE / "JalDrishti_Project_Report.docx"
PAGES = HERE / "pages.json"
FONT = "Times New Roman"

REFS = json.loads((HERE / "references.json").read_text(encoding="utf-8"))
cite_order: list[str] = []


# ── low-level helpers ────────────────────────────────────────────────
def set_font(run, size=12, bold=None, italic=None, underline=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def para_fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, spacing=1.5,
             left=None, hanging=None, keep_next=False):
    f = p.paragraph_format
    p.alignment = align
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = spacing
    if left is not None:
        f.left_indent = Cm(left)
    if hanging is not None:
        f.first_line_indent = Cm(-hanging)
    if keep_next:
        f.keep_with_next = True


def cite(m):
    keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
    nums = []
    for k in keys:
        if k not in REFS:
            raise KeyError(f"unknown reference key {k!r}")
        if k not in cite_order:
            cite_order.append(k)
        nums.append(cite_order.index(k) + 1)
    return ", ".join(f"[{n}]" for n in nums)


def add_rich(p, text, size=12, bold=False, italic=False):
    """Add text with **bold**, *italic* and [@key] citations to paragraph p."""
    text = re.sub(r"\[(@[^\]]+)\]", cite, text)
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); set_font(r, size, bold=True, italic=italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); set_font(r, size, bold=bold, italic=True)
        else:
            r = p.add_run(part); set_font(r, size, bold=bold, italic=italic)


def bottom_border(p, sz=8):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "1"), ("w:color", "000000")):
        e.set(qn(k), v)
    b.append(e)
    pPr.append(b)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def cell_borders(cell, on=True):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if on else "nil")
        if on:
            e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000")
        borders.append(e)
    tcPr.append(borders)


def shade(cell, fill="E7E6E6"):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill)
    tcPr.append(s)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true")
    trPr.append(h)


def cell_text(cell, text, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    para_fmt(p, align=align, after=0, spacing=1.0)
    add_rich(p, text, size=size, bold=bold)


def add_table(doc, rows, widths_cm=None, header=True, grid=True, size=11, center_cols=()):
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths_cm is None:
        total = 16.0
        lens = [max(len(re.sub(r"\*", "", r[c])) for r in rows) for c in range(ncol)]
        lens = [min(max(l, 6), 60) for l in lens]
        widths_cm = [total * l / sum(lens) for l in lens]
    for i, r in enumerate(rows):
        row = t.rows[i]
        for c in range(ncol):
            cell = row.cells[c]
            cell.width = Cm(widths_cm[c])
            align = WD_ALIGN_PARAGRAPH.CENTER if c in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cell, r[c].strip(), size=size, bold=(header and i == 0), align=align)
            cell_borders(cell, grid)
            if header and i == 0 and grid:
                shade(cell)
        if header and i == 0:
            repeat_header(row)
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        if len(rows) <= 18 and i < len(rows) - 1:      # keep short tables on one page
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ── document setup ───────────────────────────────────────────────────
def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(12)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        h = doc.styles[name]
        h.font.name = FONT
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
    # page number, centred, in the footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt
            r._r.append(it)
    set_font(r, 12)
    return doc


# ── block renderers ──────────────────────────────────────────────────
class Counter:
    chapter = 0
    fig = 0
    tab = 0


def chapter_heading(doc, text, first):
    if not first:
        page_break(doc)
    m = re.match(r"(\d+)\.", text)
    if m:
        Counter.chapter, Counter.fig, Counter.tab = int(m.group(1)), 0, 0
    p = doc.add_paragraph(style="Heading 1")
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=12, spacing=1.15, keep_next=True)
    add_rich(p, text, size=14, bold=True)


def section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=6, spacing=1.15, keep_next=True)
    add_rich(p, text, size=12, bold=True)


def minor_heading(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=4, spacing=1.15, keep_next=True)
    add_rich(p, text, size=12, bold=True)


def paragraph(doc, text):
    p = doc.add_paragraph()
    para_fmt(p)
    add_rich(p, text)


def bullet(doc, text, level=1):
    p = doc.add_paragraph(style="List Bullet" if level == 1 else "List Bullet 2")
    para_fmt(p, after=2, spacing=1.3, left=1.0 if level == 1 else 1.8, hanging=0.5)
    add_rich(p, text)


def numbered(doc, num, text):
    p = doc.add_paragraph()
    para_fmt(p, after=3, spacing=1.4, left=1.0, hanging=0.7)
    add_rich(p, f"{num}\t{text}")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))


def figure(doc, caption, path, width_cm):
    Counter.fig += 1
    src = (HERE / path) if not Path(path).is_absolute() else Path(path)
    if not src.exists():
        src = REPO / path
    p = doc.add_paragraph()
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, spacing=1.0, keep_next=True)
    p.add_run().add_picture(str(src), width=Cm(width_cm))
    c = doc.add_paragraph()
    para_fmt(c, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, spacing=1.0)
    add_rich(c, f"Figure {Counter.chapter}.{Counter.fig}: {caption}", size=11, italic=True)


def table_block(doc, caption, rows):
    Counter.tab += 1
    c = doc.add_paragraph()
    para_fmt(c, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4, spacing=1.0, keep_next=True)
    add_rich(c, f"Table {Counter.chapter}.{Counter.tab}: {caption}", size=11, bold=True)
    widths = None
    if rows and rows[0] and rows[0][0].startswith("{w="):
        widths = [float(x) for x in rows[0][0][3:-1].split(",")]
        rows = rows[1:]
    add_table(doc, rows, widths_cm=widths)


def flow(doc, steps):
    for i, s in enumerate(steps):
        p = doc.add_paragraph()
        para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, spacing=1.2, keep_next=i < len(steps) - 1)
        add_rich(p, s.strip(), bold=True)
        if i < len(steps) - 1:
            a = doc.add_paragraph()
            para_fmt(a, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, spacing=1.0, keep_next=True)
            set_font(a.add_run("↓"), 14)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def equation(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=8, spacing=1.2)
    add_rich(p, text, italic=True)


def render_markup(doc, text, first_chapter=True):
    blocks = re.split(r"\n\s*\n", text.strip())
    first = first_chapter
    for block in blocks:
        lines = [l.rstrip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue
        head = lines[0]
        if head.startswith("# "):
            chapter_heading(doc, head[2:].strip(), first); first = False
            rest = "\n".join(lines[1:])
            if rest:
                render_markup(doc, rest, first_chapter=False)
        elif head.startswith("## "):
            section_heading(doc, head[3:].strip())
            if len(lines) > 1:
                render_markup(doc, "\n".join(lines[1:]), first_chapter=False)
        elif head.startswith("### "):
            minor_heading(doc, head[4:].strip())
            if len(lines) > 1:
                render_markup(doc, "\n".join(lines[1:]), first_chapter=False)
        elif head == "PAGEBREAK":
            page_break(doc)
        elif head.startswith("TABLE:"):
            rows = [[c for c in l.strip().strip("|").split("|")] for l in lines[1:]]
            table_block(doc, head[6:].strip(), rows)
        elif head.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*?)\)(\{w=([\d.]+)\})?", head)
            figure(doc, m.group(1), m.group(2), float(m.group(4) or 15))
        elif head.startswith("FLOW:"):
            flow(doc, head[5:].split(">>"))
        elif head.startswith("EQ:"):
            for l in lines:
                equation(doc, l[3:].strip())
        elif head.startswith("- ") or head.startswith("-- ") or re.match(r"\d+\.\s", head):
            for l in lines:
                if l.startswith("-- "):
                    bullet(doc, l[3:], level=2)
                elif l.startswith("- "):
                    bullet(doc, l[2:])
                else:
                    m = re.match(r"(\d+\.)\s+(.*)", l)
                    if m:
                        numbered(doc, m.group(1), m.group(2))
                    else:  # continuation line of the previous item
                        doc.paragraphs[-1].add_run(" ")
                        add_rich(doc.paragraphs[-1], l.strip())
        else:
            paragraph(doc, " ".join(l.strip() for l in lines))


def code_block(doc, code):
    lines = code.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=4 if i == 0 else 0,
                 after=6 if i == len(lines) - 1 else 0, spacing=1.0, left=0.3)
        pPr = p._p.get_or_add_pPr()
        s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), "F2F2F2")
        pPr.append(s)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        r.font.size = Pt(8.5)


def render_text(doc, text, first_chapter):
    """Render markup, handling ``` fenced code blocks (which may contain blank lines)."""
    parts = re.split(r"^```[a-z]*\s*$", text, flags=re.M)
    for i, part in enumerate(parts):
        if i % 2 == 1:
            code_block(doc, part)
        elif part.strip():
            render_markup(doc, part, first_chapter=first_chapter)
            first_chapter = False


def fixed_table(t, widths_cm):
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(int(sum(widths_cm) * 567))); w.set(qn("w:type"), "dxa")
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(w)
    for row in t.rows:
        for c, width in zip(row.cells, widths_cm):
            c.width = Cm(width)


# ── front and back matter ────────────────────────────────────────────
def centred(doc, text, size=12, bold=False, italic=False, after=6, before=0, underline=False):
    p = doc.add_paragraph()
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=before, after=after, spacing=1.15)
    add_rich(p, text, size=size, bold=bold, italic=italic)
    if underline:
        for r in p.runs:
            r.underline = True
    return p


def cover(doc, F):
    for _ in range(3):
        doc.add_paragraph()
    centred(doc, "A PROJECT REPORT", 14, bold=True, after=4)
    centred(doc, "ON", 13, bold=True, after=8)
    centred(doc, F["title"], 16, bold=True, after=4)
    centred(doc, F["subtitle"], 13, bold=True, after=10)
    centred(doc, "Submitted to", 12, italic=True, after=6)
    p = doc.add_paragraph(); para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, spacing=1.0)
    p.add_run().add_picture(str(HERE / "assets/texmin_logo.png"), height=Cm(2.0))
    p.add_run("      ")
    p.add_run().add_picture(str(HERE / "assets/bit_sindri_crest.png"), height=Cm(2.3))
    centred(doc, "TEXMiN-BIT Sindri CPS CoE", 13, bold=True, after=4, before=6)
    centred(doc, "Submitted in Partial Fulfilment of the requirements for", 12, after=4)
    centred(doc, "UDAAN UG Fellowship Programme (2025-26)", 13, bold=True, after=28)
    centred(doc, "Submitted By", 12, after=4)
    centred(doc, F["student_upper"], 13, bold=True, after=4)
    centred(doc, f"(Roll No. {F['roll']})", 12, bold=True, after=4)
    centred(doc, F["student_dept"], 12, bold=True, after=28)
    centred(doc, "Under the Guidance of", 12, bold=True, after=4)
    centred(doc, F["mentor_upper"], 13, bold=True, after=4)
    centred(doc, F["mentor_dept"], 12, bold=True, after=4)


def heading_line(doc, text, size=14, underline_rule=True, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    para_fmt(p, align=align, before=24, after=14, spacing=1.15)
    add_rich(p, text, size=size, bold=True)
    if underline_rule:
        bottom_border(p, 6)


def signature_pair(doc, left, right, before=36):
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    fixed_table(t, [8.2, 7.3])
    for i, lines in enumerate((left, right)):
        c = t.rows[0].cells[i]
        c.width = Cm(8.0)
        cell_borders(c, False)
        c.text = ""
        for j, l in enumerate(lines):
            p = c.paragraphs[0] if j == 0 else c.add_paragraph()
            para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.LEFT,
                     before=before if j == 0 else 0, after=2, spacing=1.15)
            add_rich(p, l, size=12)


def signature_centre(doc, lines, before=30):
    for j, l in enumerate(lines):
        p = doc.add_paragraph()
        para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=before if j == 0 else 0, after=2, spacing=1.15)
        add_rich(p, l)


def contents(doc, entries, pages):
    heading_line(doc, "CONTENTS:", size=14, align=WD_ALIGN_PARAGRAPH.LEFT)
    rows = [["Chapter", "Title", "Page no."]]
    for chap, title, key in entries:
        rows.append([chap, title, pages.get(key, "")])
    t = add_table(doc, rows, widths_cm=[3.2, 9.8, 2.6], size=12, center_cols=(0, 2))
    for r in t.rows[1:]:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.bold = True
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)


def front_matter(doc, F, pages):
    cover(doc, F)
    page_break(doc)
    # certificate
    heading_line(doc, "CERTIFICATE")
    render_markup(doc, F["certificate"])
    signature_pair(doc,
                   ["(Signature of Centre in Charge)", "Date:", F["centre_in_charge"], "TEXMiN BIT Sindri CPS CoE"],
                   ["(Signature of Supervisor)", "Date:", F["mentor"], F["mentor_dept_short"]])
    signature_centre(doc, ["(Signature of Project Coordinator)", "Date:", F["coordinator"], "TEXMiN-BIT Sindri CPS CoE"])
    page_break(doc)
    heading_line(doc, "CANDIDATE’S DECLARATION")
    render_markup(doc, F["declaration"])
    signature_pair(doc, ["(Signature of Student)", "", F["student"], f"Roll No: {F['roll']}",
                         "Department of Information Technology", "B.Tech, Session 2024-28", "B.I.T. Sindri"],
                   ["Place: Sindri, Jharkhand", "Date:"])
    page_break(doc)
    heading_line(doc, "ACKNOWLEDGEMENT", underline_rule=False)
    render_markup(doc, F["acknowledgement"])
    p = doc.add_paragraph(); para_fmt(p, align=WD_ALIGN_PARAGRAPH.RIGHT, before=24, after=0)
    add_rich(p, F["student"], bold=True)
    page_break(doc)
    heading_line(doc, "ABSTRACT", underline_rule=False)
    n0 = len(doc.paragraphs)
    render_markup(doc, F["abstract"])
    for p in doc.paragraphs[n0:]:
        p.paragraph_format.line_spacing = 1.35
    p = doc.add_paragraph(); para_fmt(p, before=8)
    add_rich(p, "**Keywords:** " + F["keywords"])
    page_break(doc)
    contents(doc, F["contents"], pages)


def back_matter(doc, B):
    page_break(doc)
    heading_line(doc, "MENTOR’S REMARKS & RECOMMENDATION", align=WD_ALIGN_PARAGRAPH.LEFT, underline_rule=False)
    render_markup(doc, B["remarks"])
    for line in ("**Remarks:** ____________________________________________________",
                 "_________________________________________________________________",
                 "**Recommendation:**   ☐ Satisfactory      ☐ Needs Improvement      ☐ Excellent",
                 "**Signature of Mentor:** ____________________________",
                 f"**{B['mentor']}**",
                 "**Date:** _______________________"):
        p = doc.add_paragraph(); para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4)
        add_rich(p, line)
    page_break(doc)
    heading_line(doc, "CERTIFICATE OF APPROVAL", underline_rule=False)
    doc.paragraphs[-1].runs[0].underline = True
    render_markup(doc, B["approval"])
    signature_pair(doc,
                   ["**" + B["mentor"] + "**", "Mentor", "Department of Production and Industrial Engineering",
                    "BIT Sindri, Dhanbad", "Date: ____________________"],
                   ["**" + B["coordinator"] + "**", "Project Coordinator,", "TEXMiN-BIT Sindri CPS CoE",
                    "BIT Sindri, Dhanbad", "Date: ____________________"], before=60)
    page_break(doc)
    heading_line(doc, "SUBMISSION CHECKLIST", align=WD_ALIGN_PARAGRAPH.LEFT, underline_rule=False)
    paragraph(doc, "The following documents and materials have been submitted along with the fellowship report:")
    rows = [["Sl. No.", "Item Submitted", "Status (✓/✗)"]]
    for item in B["checklist"]:
        if item.startswith("- "):
            rows.append(["", item[2:], "☐"])
        else:
            n, txt = item.split(" ", 1)
            rows.append([n, f"**{txt}**", "☐"])
    t = add_table(doc, rows, widths_cm=[2.0, 11.0, 3.0], grid=False, size=12, center_cols=(0, 2))
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)


def references(doc):
    for i, key in enumerate(cite_order, start=1):
        p = doc.add_paragraph()
        para_fmt(p, after=4, spacing=1.2, left=1.0, hanging=1.0)
        add_rich(p, f"[{i}]\t{REFS[key]}")
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))


# ── build ────────────────────────────────────────────────────────────
def build(pages):
    cite_order.clear()
    Counter.chapter = Counter.fig = Counter.tab = 0
    F = json.loads((HERE / "front.json").read_text(encoding="utf-8"))
    B = json.loads((HERE / "back.json").read_text(encoding="utf-8"))
    doc = new_document()
    front_matter(doc, F, pages)
    chapters = sorted(HERE.glob("ch*.txt"))
    first = False
    for f in chapters:
        text = f.read_text(encoding="utf-8")
        if "{{REFERENCES}}" in text:
            before, after = text.split("{{REFERENCES}}")
            render_text(doc, before, first_chapter=first)
            references(doc)
            if after.strip():
                render_text(doc, after, first_chapter=False)
        else:
            render_text(doc, text, first_chapter=first)
    back_matter(doc, B)
    doc.save(OUT)
    unused = sorted(set(REFS) - set(cite_order))
    print(f"built {OUT.name}: {len(cite_order)} references cited" + (f"; unused keys: {unused}" if unused else ""))


def render_pdf():
    pdf = OUT.with_suffix(".pdf")
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{OUT}'); $d.Fields.Update() | Out-Null; "
          f"$d.SaveAs2('{pdf}', 17); $d.Close(0); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    return pdf


def find_pages(pdf):
    import fitz
    F = json.loads((HERE / "front.json").read_text(encoding="utf-8"))
    d = fitz.open(pdf)
    texts = [pg.get_text() for pg in d]
    starts = {}
    contents_page = next(i for i, t in enumerate(texts) if "CONTENTS:" in t)
    for chap, title, key in F["contents"]:
        marker = F["markers"][key]
        lo = 0 if key in ("certificate", "declaration", "acknowledgement", "abstract") else contents_page + 1
        for i in range(lo, len(texts)):
            first_lines = "\n".join(texts[i].strip().split("\n")[:4])
            if marker in first_lines:
                starts[key] = i + 1
                break
    keys = [k for _, _, k in F["contents"]]
    bounds = sorted(set(starts.values()) | {contents_page + 1})   # the contents page is a boundary too
    pages = {}
    for k in keys:
        s = starts.get(k)
        if s is None:
            continue
        later = [b for b in bounds if b > s]
        e = (later[0] - 1) if later else s
        pages[k] = str(s) if e <= s else f"{s}-{e}"
    missing = [k for k in keys if k not in starts]
    if missing:
        print("page markers not found:", missing)
    return pages, len(d)


if __name__ == "__main__":
    pages = json.loads(PAGES.read_text()) if PAGES.exists() else {}
    build(pages)
    if "--pages" in sys.argv:
        for _ in range(2):
            pdf = render_pdf()
            pages, n = find_pages(pdf)
            PAGES.write_text(json.dumps(pages, indent=1))
            build(pages)
        pdf = render_pdf()
        print(f"{n} pages; contents: {pages}")
